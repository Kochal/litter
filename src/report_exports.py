"""Exports for the story report (docs/story.md): the data behind every chart as an
Excel workbook, one sheet per chart, and the drive-time map as GeoParquet.

Inputs: outputs/report/story_data.json (the numbers the story page draws),
outputs/rise_factors_by_year.csv, outputs/composition_change.csv,
the drive-time panel (drive_time_map.py), the verified recycling centre history and LSOA
boundaries (ONS, super generalised).
Outputs (outputs/report/):
- fly_tipping_story_data.xlsx: a contents sheet, then one sheet per chart;
  derived columns (percentage changes, indexes, rates) are Excel formulas.
- drive_time_change_lsoa.parquet: GeoParquet, one row per neighbourhood (LSOA
  2021), drive time to the nearest recycling centre in 2012 and 2024 and the
  change, with boundaries in British National Grid (EPSG:27700).
- recycling_centres.parquet: GeoParquet points, every English centre with the years it
  appears open in the verified history, and whether it closed between 2012 and
  2023 (the crosses on the map).
- lsoa_population_weighted_centroids.parquet: GeoParquet points, the ONS
  population-weighted centroid of every neighbourhood (LSOA 2021), the point
  drive times are measured from, with its council, rural-urban class, residents,
  and drive times in 2012 and 2024. 8 of 35,672 have no drive time because they
  could not be routed on the road network: the Isles of Scilly (no road link) and
  7 whose centroid lies on a road not connected to the main OS Open Roads network.
No personal data: only published aggregates and public facility names.
"""
import json
from pathlib import Path

import geopandas as gpd
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
REP = OUT / "report"
FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="E2EEE9")
BANDS = [(-99, -1, "Shorter by 1+ min"), (-1, 1, "Little change (under 1 min)"), (1, 3, "Longer by 1 to 3 min"),
         (3, 5, "Longer by 3 to 5 min"), (5, 99, "Longer by 5+ min")]


def fy(y: int) -> str:
    return f"{y}/{str(y + 1)[2:]}"


class Sheet:
    """Writes a titled sheet: title, note on observations, header row, rows."""

    def __init__(self, wb, name, title, note):
        self.ws = wb.create_sheet(name)
        self.ws["A1"] = title
        self.ws["A1"].font = Font(name=FONT, bold=True, size=13)
        self.ws["A2"] = note
        self.ws["A2"].font = Font(name=FONT, italic=True, size=10)
        self.ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        self.ws.row_dimensions[2].height = 45
        self.row = 4

    def header(self, cols, widths):
        for j, (c, w) in enumerate(zip(cols, widths), 1):
            cell = self.ws.cell(self.row, j, c)
            cell.font = Font(name=FONT, bold=True)
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            self.ws.column_dimensions[get_column_letter(j)].width = w
        self.ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(cols), 4))
        self.ws.freeze_panes = self.ws.cell(self.row + 1, 1)
        self.first = self.row + 1
        self.row += 1

    def add(self, values, formats=None):
        """values may contain callables r -> formula string, given the row number."""
        for j, v in enumerate(values, 1):
            cell = self.ws.cell(self.row, j, v(self.row) if callable(v) else v)
            cell.font = Font(name=FONT)
            if formats and formats.get(j):
                cell.number_format = formats[j]
        self.row += 1

    def footnote(self, text):
        self.row += 1
        c = self.ws.cell(self.row, 1, text)
        c.font = Font(name=FONT, italic=True, size=9)
        self.row += 1


def ratio_cols(r, start):
    """Rate ratio, low, high, then their percentage changes as formulas."""
    a, b, c = (get_column_letter(start + i) for i in range(3))
    return [r["irr"], r["lo"], r["hi"], lambda n: f"={a}{n}-1", lambda n: f"={b}{n}-1", lambda n: f"={c}{n}-1"]


RATIO_FMT = {"r": "0.000", "p": "+0%;-0%;0%"}


def ratio_formats(start):
    return {start + i: ("0.000" if i < 3 else "+0%;-0%;0%") for i in range(6)}


def workbook(d: dict) -> Workbook:
    wb = Workbook()
    contents = wb.active
    contents.title = "Contents"
    sheets = []

    # 1. Trend
    s = Sheet(wb, "1 Trend", "Figure 1. Recorded fly-tipping per 1,000 people by size of load, England",
              f"Official council counts (Defra). Size-of-load columns: the {d['rc']['steady rise'] + d['rc']['little change']} "
              "councils whose recorded numbers changed only gradually (no sudden jump). Last column: all 296 English "
              "councils. Index columns are formulas (2012/13 = 100). From 2019/20 councils were also asked to count "
              "incidents their own crews find.")
    cols = ["Year", "Small items", "Van loads", "Tipper-lorry load or larger", "All fly-tipping (no sudden jump)",
            "All fly-tipping, all 296 councils"]
    s.header(cols + [f"Index: {c}" for c in cols[1:]], [10] + [16] * 10)
    keys = ["ns_small", "ns_van", "ns_large", "ns_total", "all_total"]
    for r in d["trend"]:
        idx = [(lambda L: (lambda n: f"={L}{n}/{L}${s.first}*100"))(get_column_letter(j)) for j in range(2, 7)]
        s.add([fy(r["year"])] + [r[k] for k in keys] + idx, {**{j: "0.00" for j in range(2, 7)}, **{j: "0" for j in range(7, 12)}})
    sheets.append(s)

    # 2. Access tests
    s = Sheet(wb, "2 Access tests", "Figure 3. Every test of access to recycling centres",
              "Change in fly-tipping per resident (rate ratio, with 95% likely range). The observations column gives "
              "the number of councils, council-years, neighbourhoods or closures behind each test.")
    s.header(["Test", "Observations", "Rate ratio", "Low", "High", "Change", "Change, low", "Change, high"],
             [58, 52, 11, 9, 9, 10, 11, 11])
    for label, sub, r in d["access"]:
        s.add([label, sub] + ratio_cols(r, 3), ratio_formats(3))
    e = d["excep"]
    s.add(["Exception: a closure added 3 minutes or more, years after the closure",
           f"{e['n']:,} neighbourhoods"] + ratio_cols(e["r"], 3), ratio_formats(3))
    sheets.append(s)

    # 3. Renting and car ownership
    s = Sheet(wb, "3 Renting and no car", "Figure 4. Private renting and car ownership by kind of fly-tipping",
              "Change in fly-tipping per 10 percentage points more households renting privately, or without a car or "
              "van, comparing councils with similar deprivation, density and access. 950 council-years from 317 "
              "councils, 2022/23 to 2024/25, all characteristics in one model.")
    s.header(["Kind of fly-tipping", "Renting: rate ratio", "Low", "High", "Renting: change", "Low", "High",
              "No car: rate ratio", "Low", "High", "No car: change", "Low", "High"], [44] + [11] * 12)
    for x in d["drivers"]:
        s.add([x["label"]] + ratio_cols(x["rent"], 2) + ratio_cols(x["car"], 8), {**ratio_formats(2), **ratio_formats(8)})
    sheets.append(s)

    # 4a. Land
    s = Sheet(wb, "4a Type of land", "Figure 5. Recorded fly-tipping by type of land",
              f"Official council counts, England and Wales, 2022/23 to 2024/25 ({d['land_incidents']:,} incidents). "
              "Street-type land: highways, footpaths and bridleways, council land and back alleys.")
    s.header(["Type of land", "Share of incidents", "Street-type land"], [34, 16, 16])
    street = {"Highways", "Footpaths and bridleways", "Council land", "Back alleys"}
    for x in d["land"]:
        s.add([x["label"], x["pct"] / 100, "yes" if x["label"] in street else "no"], {2: "0.0%"})
    last = s.row - 1
    s.add(["Street-type land, total", lambda n: f'=SUMIFS(B{s.first}:B{last},C{s.first}:C{last},"yes")', ""], {2: "0.0%"})
    sheets.append(s)

    # 4b. Setting
    s = Sheet(wb, "4b Street or out of way", "Figure 6. Where FixMyStreet reports say the waste was",
              "FixMyStreet reports (2012 to 2025) whose text says where the waste was: a doorstep, pavement or street, "
              "or an out-of-the-way place such as a layby, field or verge. Totals and shares are formulas.")
    s.header(["Reports", "Doorstep, pavement or street", "Out-of-the-way place", "Total", "Share in the street"],
             [30, 18, 18, 12, 14])
    for x in d["setting"]:
        s.add([x["label"], x["door"], x["out"], lambda n: f"=B{n}+C{n}", lambda n: f"=B{n}/D{n}"],
              {2: "#,##0", 3: "#,##0", 4: "#,##0", 5: "0%"})
    sheets.append(s)

    # 5a. Predicted vs actual
    w = d["why"]
    s = Sheet(wb, "5a Predicted vs actual", "Figure 7. Change predicted from renting and car ownership, against the actual change",
              "England, 296 councils. Actual: recorded fly-tipping per person, 2012/13 to 2014/15 against 2022/23 to "
              "2024/25. Predicted: national change in each Census share (2011 to 2021) times the council comparisons "
              "in figure 4 (all fly-tipping and tipper loads: whole-council model; small items and van loads: "
              "size-of-load models).")
    s.header(["Size of load", "From more renting", "From fewer car-less households", "Predicted, both", "Actual"],
             [30, 16, 18, 16, 12])
    for x in w["pred"]:
        s.add([x["label"], x["rent"], x["nocar"], lambda n: f"=(1+B{n})*(1+C{n})-1", x["actual"]],
              {j: "+0.0%;-0.0%;0.0%" for j in range(2, 6)})
    sheets.append(s)

    # 5b. Census shares
    s = Sheet(wb, "5b Census shares", "Table: share of households renting privately and without a car, 2011 and 2021",
              f"Census 2011 (KS402EW, KS404EW) and Census 2021, {w['n_shares']} councils on 2025 boundaries "
              f"({w['n_shares_e']} in England), weighted by households. Change is a formula, in percentage points.")
    s.header(["Nation", "Measure", "2011", "2021", "Change, percentage points"], [12, 26, 10, 10, 16])
    for n in ("England", "Wales"):
        for k, lab in (("rent", "Renting privately"), ("nocar", "No car or van")):
            s.add([n, lab, w["shares"][n]["2011"][k], w["shares"][n]["2021"][k], lambda r: f"=(D{r}-C{r})*100"],
                  {3: "0.0%", 4: "0.0%", 5: "+0.0;-0.0"})
    sheets.append(s)

    # 5c. Other factors by year
    y = pd.read_csv(OUT / "rise_factors_by_year.csv")
    s = Sheet(wb, "5c Other factors by year", "Figure 8. Moving, street-cleaning spending and enforcement, England",
              "296 English councils. Rates per 1,000 incidents are formulas from the counts. Moving rate: moves into "
              "and out of each council area (from the UK or abroad) per 100 residents, weighted by population (ONS). "
              "Street cleaning: net current spending (revenue outturn RO5), from 2017/18; 2024/25 prices use the ONS "
              "consumer prices index (D7BT).")
    s.header(["Year", "Councils", "Population", "Incidents", "Prosecutions", "Fixed penalty notices",
              "Prosecutions per 1,000 incidents", "Fixed penalty notices per 1,000 incidents",
              "Residents moving in or out, per 100", "Street cleaning per person, GBP, prices of each year",
              "Street cleaning per person, GBP, 2024/25 prices"], [10, 10, 13, 12, 12, 13, 15, 15, 15, 17, 17])
    for r in y.itertuples():
        s.add([fy(r.year), r.councils, r.population, r.total, r.prosecutions, r.fpn,
               lambda n: f"=E{n}/D{n}*1000", lambda n: f"=F{n}/D{n}*1000", r.churn_pct,
               None if pd.isna(r.cleansing_per_person_gbp) else r.cleansing_per_person_gbp,
               None if pd.isna(r.cleansing_per_person_gbp_2024_prices) else r.cleansing_per_person_gbp_2024_prices],
              {3: "#,##0", 4: "#,##0", 5: "#,##0", 6: "#,##0", 7: "0.00", 8: "0.0", 9: "0.0", 10: "0.00", 11: "0.00"})
    sheets.append(s)

    # 5d. Same council over time
    s = Sheet(wb, "5d Same council over time", "Figure 9. Moving rate and street-cleaning spending, each council compared with itself",
              "Poisson models with council and year fixed effects. Moving: per 5 more residents in every 100 moving in "
              "or out, 2012/13 to 2024/25. Spending: street-cleaning spending per person halved, 2017/18 to 2024/25 "
              "(model also includes the moving rate). The story shows the councils with no sudden jump in recording.")
    s.header(["Factor", "Size of load", "Councils", "Rate ratio", "Low", "High", "Change", "Change, low",
              "Change, high", "p", "Number of councils", "Council-years"], [30, 26, 30, 11, 9, 9, 10, 11, 11, 8, 11, 12])
    fac = {"churn": "5 more movers per 100 residents", "spend": "Street-cleaning spending halved"}
    for x in w["within"]:
        for smp, lab in (("ns", "No sudden jump in recording"), ("all", "All councils")):
            r = x[smp]
            s.add([fac[x["kind"]], x["label"], lab] + ratio_cols(r, 4) + [r["p"], r["nc"], r["ncy"]],
                  {**ratio_formats(4), 10: "0.000", 11: "#,##0", 12: "#,##0"})
    g = w["gap_tipper_nocar"]
    s.footnote(f"Widening gap, tipper-lorry loads: per 10 percentage points more car-less households, growth per decade "
               f"x{g['all']['irr']:.2f} (likely range {g['all']['lo']:.2f} to {g['all']['hi']:.2f}; "
               f"{g['all']['nc']} councils, {g['all']['ncy']:,} council-years); in councils with no sudden jump "
               f"x{g['ns']['irr']:.2f} ({g['ns']['lo']:.2f} to {g['ns']['hi']:.2f}; {g['ns']['nc']} councils).")
    sheets.append(s)

    # Contents
    contents["A1"] = "What drives fly-tipping in England and Wales: data behind the charts"
    contents["A1"].font = Font(name=FONT, bold=True, size=14)
    lines = [
        "One sheet per chart in the report (docs/story.md). Derived columns (percentage changes, indexes, rates, "
        "totals) are formulas. A rate ratio of 1.20 means 20% more fly-tipping per resident.",
        "The map of drive-time changes (figure 2) is in drive_time_change_lsoa.parquet and recycling_centres.parquet "
        "(GeoParquet, British National Grid; centres for England only, as on the map), which open in QGIS, ArcGIS, R (sf) or Python (geopandas). "
        "lsoa_population_weighted_centroids.parquet has the point drive times are measured from for every neighbourhood.",
        "Sources: Defra and Welsh Government fly-tipping statistics; FixMyStreet (mySociety); ONS Census 2011 and 2021, "
        "population estimates and consumer prices index; council revenue outturn (MHCLG); Environment Agency and "
        "Natural Resources Wales site records. Open Government Licence v3.0. Code: github.com/kochal/litter.",
    ]
    for i, t in enumerate(lines, 3):
        c = contents.cell(i, 1, t)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        contents.row_dimensions[i].height = 48
    contents.merge_cells("A3:C3"), contents.merge_cells("A4:C4"), contents.merge_cells("A5:C5")
    for j, (h, wdt) in enumerate((("Sheet", 26), ("Chart", 70), ("Rows", 8)), 1):
        c = contents.cell(7, j, h)
        c.font, c.fill = Font(name=FONT, bold=True), HEAD_FILL
        contents.column_dimensions[get_column_letter(j)].width = wdt
    for i, sh in enumerate(sheets, 8):
        c = contents.cell(i, 1, sh.ws.title)
        c.hyperlink = f"#'{sh.ws.title}'!A1"
        c.font = Font(name=FONT, color="1D5B4C", underline="single")
        contents.cell(i, 2, sh.ws["A1"].value).font = Font(name=FONT)
        contents.cell(i, 3, sh.row - sh.first).font = Font(name=FONT)
    return wb


def map_layers():
    from drive_time_map import change   # unrounded, as on the map
    c = change()
    c["band"] = pd.cut(c["change"], [b[0] for b in BANDS] + [99], labels=[b[2] for b in BANDS], right=False).astype(str)
    lsoa = gpd.read_file(ROOT / "data" / "raw" / "lsoa21_bsc_ruc.gpkg", columns=["LSOA21CD"])
    g = lsoa.merge(c, on="LSOA21CD").rename(columns={
        "t2012": "drive_min_2012", "t2024": "drive_min_2024", "change": "change_min_2012_2024",
        "max_rise": "largest_rise_min"})
    g["nation"] = g["LSOA21CD"].str[0].map({"E": "England", "W": "Wales"})
    g = g.to_crs(27700)
    g.to_parquet(REP / "drive_time_change_lsoa.parquet", index=False)
    h = pd.read_csv(ROOT / "data" / "interim" / "hwrc_england_history_verified.csv")
    h["closed_2012_2023"] = h["last"].between(2012, 2023)
    pts = gpd.GeoDataFrame(h.rename(columns={"first": "first_year_open", "last": "last_year_open"}),
                           geometry=gpd.points_from_xy(h["easting"], h["northing"]), crs=27700)
    pts.to_parquet(REP / "recycling_centres.parquet", index=False)
    return len(g), int((g["change_min_2012_2024"] >= 3).sum()), int(pts["closed_2012_2023"].sum())


def centroids():
    from access_panel import lsoa_to_lad
    from drive_time_map import change
    raw = ROOT / "data" / "raw"
    pwc = gpd.read_file(raw / "lsoa21_pwc.gpkg", columns=["LSOA21CD"])
    attrs = gpd.read_file(raw / "lsoa21_bsc_ruc.gpkg", columns=["LSOA21CD", "LSOA21NM", "RUC21CD", "RUC21NM"],
                          ignore_geometry=True)
    lad = gpd.read_file(raw / "lad25_bgc.gpkg", columns=["LAD25CD", "LAD25NM"], ignore_geometry=True)
    cen = pd.read_csv(ROOT / "data" / "interim" / "census_lsoa.csv")[["LSOA21CD", "residents", "households"]]
    t = change().rename(columns={"t2012": "drive_min_2012", "t2024": "drive_min_2024",
                                 "change": "change_min_2012_2024", "max_rise": "largest_rise_min"})
    g = (pwc.merge(attrs, on="LSOA21CD", how="left").merge(lsoa_to_lad(), on="LSOA21CD", how="left")
         .merge(lad, on="LAD25CD", how="left").merge(cen, on="LSOA21CD", how="left").merge(t, on="LSOA21CD", how="left"))
    g = g.to_crs(27700)
    g["easting"], g["northing"] = g.geometry.x.round(1), g.geometry.y.round(1)
    ll = g.geometry.to_crs(4326)
    g["lat"], g["lon"] = ll.y.round(6), ll.x.round(6)
    g["nation"] = g["LSOA21CD"].str[0].map({"E": "England", "W": "Wales"})
    cols = ["LSOA21CD", "LSOA21NM", "LAD25CD", "LAD25NM", "nation", "RUC21CD", "RUC21NM", "residents", "households",
            "easting", "northing", "lat", "lon", "drive_min_2012", "drive_min_2024", "change_min_2012_2024",
            "largest_rise_min", "geometry"]
    g[cols].to_parquet(REP / "lsoa_population_weighted_centroids.parquet", index=False)
    return len(g), int(g["drive_min_2012"].isna().sum())


if __name__ == "__main__":
    REP.mkdir(exist_ok=True)
    d = json.loads((REP / "story_data.json").read_text())
    workbook(d).save(REP / "fly_tipping_story_data.xlsx")
    print("neighbourhoods, 3+ min longer, closures:", map_layers())
    print("centroids, without drive times:", centroids())
